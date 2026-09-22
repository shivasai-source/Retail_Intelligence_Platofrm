"""Decision Center board decisions -- the store and the export.

The board is the Simulation Studio's snapshots (display strings), the chosen
slot and the rationale. The store keeps it whole; the export prints it; nothing
recomputes a figure.

Run with:

    ../venv/Scripts/python.exe -m pytest tests/test_board_decisions.py -q
"""

from __future__ import annotations

import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader

from app.main import app
from app.store import db


@pytest.fixture(scope="module", autouse=True)
def store(tmp_path_factory):
    db.use_path(tmp_path_factory.mktemp("board") / "board.db")
    yield
    db.close()


@pytest.fixture(scope="module")
def client():
    return TestClient(app)


def _scenario(slot: int, revenue: str, roi: str) -> dict:
    return {
        "id": f"scn-{slot}", "slot": slot, "signature": f"sig-{slot}", "addedAt": 1,
        "scope": {"label": "F26 (Annual) · Modern Trade", "filters": {"year": 2026, "channel": ["CH002"]}, "currency": "INR"},
        "levers": {"discount_pct": 10 + slot, "trade_spend": 1e7, "days": 14},
        "kpis": [
            {"key": "discount", "label": "Discount", "value": f"{10 + slot}.0%", "raw": 10 + slot},
            {"key": "days", "label": "Days", "value": "14 days"},
            {"key": "budget", "label": "Trade spend budget", "value": "₹1.0 Cr", "sub": "covers the whole scope"},
            {"key": "revenue", "label": "Revenue", "value": revenue, "sub": "+₹1.0 Cr (+2.0%) vs current plan"},
            {"key": "roi", "label": "ROI", "value": roi, "tone": "success"},
        ],
    }


BOARD = {"scenarios": [_scenario(1, "₹50.0 Cr", "1.40"), _scenario(2, "₹52.0 Cr", "1.31")],
         "chosen_slot": 2, "rationale": "Higher revenue at an ROI still well above break-even."}


def test_a_board_is_stored_whole_and_read_back(client):
    saved = client.post("/api/store/board-decisions", json=BOARD)
    assert saved.status_code == 200, saved.text
    body = saved.json()
    assert body["decision_id"].startswith("dec_")
    assert body["chosen_slot"] == 2
    assert body["scenarios"] == BOARD["scenarios"]
    assert body["rationale"] == BOARD["rationale"]
    assert body["dataset_version"]

    again = client.get(f"/api/store/board-decisions/{body['decision_id']}").json()
    assert again["scenarios"] == BOARD["scenarios"]


def test_the_history_lists_newest_first_with_a_summary(client):
    client.post("/api/store/board-decisions", json={**BOARD, "chosen_slot": 1, "rationale": "second"})
    listed = client.get("/api/store/board-decisions").json()["decisions"]
    assert len(listed) >= 2
    assert listed[0]["rationale"] == "second"
    assert listed[0]["chosen_slot"] == 1
    assert listed[0]["summary"]["revenue"] == "₹50.0 Cr"
    assert listed[0]["summary"]["roi"] == "1.40"
    assert listed[0]["scenario_count"] == 2


@pytest.mark.parametrize("bad,why", [
    ({"scenarios": []}, "at least one"),
    ({"scenarios": [_scenario(i, "x", "y") for i in range(1, 5)]}, "at most three"),
    ({"scenarios": [_scenario(1, "x", "y")], "chosen_slot": 3}, "not on the board"),
])
def test_a_malformed_board_is_refused(client, bad, why):
    r = client.post("/api/store/board-decisions", json=bad)
    assert r.status_code == 422, r.text
    assert why in r.json()["detail"]


def test_an_unknown_board_decision_is_404(client):
    assert client.get("/api/store/board-decisions/dec_nope").status_code == 404


def test_the_export_prints_the_board_it_was_given(client):
    r = client.post("/api/reports", json={
        "module": "decision-center", "scope": {"year": 2026, "channel": ["CH002"]},
        "options": {"board": BOARD}, "formats": ["pdf", "xlsx"],
    })
    assert r.status_code == 201, r.text
    report_id = r.json()["report_id"]
    pdf = client.get(f"/api/reports/{report_id}/download/pdf").content
    text = "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(pdf)).pages)
    assert "Scenario 2 (chosen)" in text
    assert "Higher revenue" in text
    assert "1.31" in text and "1.40" in text
    xlsx = client.get(f"/api/reports/{report_id}/download/xlsx").content
    assert "Comparison" in load_workbook(io.BytesIO(xlsx)).sheetnames


def test_the_export_refuses_an_empty_board(client):
    r = client.post("/api/reports", json={
        "module": "decision-center", "scope": {"year": 2026}, "options": {"board": {"scenarios": []}},
        "formats": ["pdf"],
    })
    assert r.status_code == 422
    assert "at least one" in r.json()["detail"]


def test_clearing_the_history_empties_it(client):
    assert client.delete("/api/store/board-decisions").json()["removed"] >= 1
    assert client.get("/api/store/board-decisions").json()["decisions"] == []
